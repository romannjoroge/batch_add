# -*- coding: utf-8 -*-
import base64
import csv
import io
import re
from difflib import SequenceMatcher

from odoo import api, fields, models
from odoo.exceptions import UserError


class StockPickingBatchImportWizard(models.TransientModel):
    _name = "stock.picking.batch.import.wizard"
    _description = "Stock Picking Batch Import Wizard"

    picking_id = fields.Many2one("stock.picking", required=True, readonly=True)
    file_data = fields.Binary(string="CSV File", required=True)
    file_name = fields.Char(string="File Name")
    sheet_name = fields.Char(
        string="Sheet",
        help="For Excel files with multiple sheets, pick the sheet to import from.",
    )
    line_ids = fields.One2many(
        "stock.picking.batch.import.line",
        "wizard_id",
        string="Lines",
    )

    def action_parse_file(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError("Please upload a CSV file first.")
        if self.file_name and not self._is_supported_file(self.file_name):
            raise UserError("Please upload a CSV or Excel file (.csv, .xls, .xlsx).")

        self.line_ids.unlink()
        rows, headers = self._parse_file()
        if not rows:
            raise UserError("No rows were found in the file.")

        qty_key = self._detect_qty_key(headers)
        descriptor_keys = self._detect_descriptor_keys(headers, qty_key)

        lines = []
        for index, row in enumerate(rows, start=1):
            descriptor = self._build_descriptor(row, headers, descriptor_keys, qty_key)
            quantity = self._extract_qty(row, headers, qty_key)
            if not descriptor:
                continue
            match = self._match_product(row, headers, descriptor)

            lines.append(
                (
                    0,
                    0,
                    {
                        "row_no": index,
                        "raw_data": self._render_row(row, headers),
                        "description": descriptor,
                        "quantity": quantity,
                        "match_status": match["status"],
                        "product_id": match.get("product_id"),
                        "candidate_product_ids": [
                            (6, 0, match.get("candidate_ids", []))
                        ],
                    },
                )
            )

        self.line_ids = lines
        return {
            "type": "ir.actions.act_window",
            "res_model": "stock.picking.batch.import.wizard",
            "res_id": self.id,
            "views": [(False, "form")],
            "target": "new",
        }

    @api.onchange("file_data", "file_name")
    def _onchange_file_data_set_sheet(self):
        for wizard in self:
            if not wizard.file_name or not wizard.file_data:
                wizard.sheet_name = False
                continue
            filename = (wizard.file_name or "").lower()
            if not (filename.endswith(".xlsx") or filename.endswith(".xls")):
                wizard.sheet_name = False
                continue

            sheet_names = wizard._get_excel_sheet_names()
            wizard.sheet_name = sheet_names[-1] if sheet_names else False

    def _get_excel_sheet_names(self):
        decoded = self._decode_binary_file()
        filename = (self.file_name or "").lower()

        if filename.endswith(".xlsx"):
            try:
                import openpyxl
            except Exception:
                return []
            wb = openpyxl.load_workbook(io.BytesIO(decoded), read_only=True, data_only=True)
            names = list(wb.sheetnames or [])
            wb.close()
            return names

        if filename.endswith(".xls"):
            if self._looks_like_html_table(decoded):
                return ["Sheet1"]
            try:
                import xlrd
            except Exception:
                return []
            wb = xlrd.open_workbook(file_contents=decoded)
            return list(getattr(wb, "sheet_names", lambda: [])())

        return []

    def action_apply_lines(self):
        self.ensure_one()
        if not self.line_ids:
            raise UserError("There are no lines to apply. Parse a CSV first.")

        lines_to_add = self.line_ids.filtered(lambda l: l.product_id and l.quantity > 0)
        if not lines_to_add:
            raise UserError("Please select at least one product to add.")

        Move = self.env["stock.move"]
        move_vals = []
        for line in lines_to_add:
            vals = {
                "picking_id": self.picking_id.id,
                "name": line.product_id.display_name,
                "product_id": line.product_id.id,
                "product_uom_qty": line.quantity,
                "product_uom": line.product_id.uom_id.id,
                "location_id": self.picking_id.location_id.id,
                "location_dest_id": self.picking_id.location_dest_id.id,
            }

            # Use a "new" record + onchanges where available to align with UI logic.
            new_move = Move.new(vals)
            if hasattr(new_move, "_onchange_product_id"):
                new_move._onchange_product_id()
            if hasattr(new_move, "_onchange_product_uom_qty"):
                new_move._onchange_product_uom_qty()
            if hasattr(new_move, "_onchange_quantity"):
                new_move._onchange_quantity()
            move_vals.append((0, 0, new_move._convert_to_write(new_move._cache)))

        self.picking_id.write({"move_ids_without_package": move_vals})
        return {"type": "ir.actions.act_window_close"}

    # -----------------------------
    # Shared parsing/matching logic
    # -----------------------------

    def _is_supported_file(self, filename):
        lower = (filename or "").lower()
        return lower.endswith(".csv") or lower.endswith(".xls") or lower.endswith(".xlsx")

    def _parse_file(self):
        filename = (self.file_name or "").lower()
        if filename.endswith(".xls") or filename.endswith(".xlsx"):
            return self._parse_excel()
        return self._parse_csv()

    def _parse_csv(self):
        decoded = self._decode_file_data()
        sample = decoded[:4096]

        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=[",", ";", "\t", "|"])
        except Exception:
            dialect = csv.get_dialect("excel")

        reader = csv.reader(io.StringIO(decoded), dialect)
        rows = list(reader)
        rows = [row for row in rows if any(col.strip() for col in row)]
        if not rows:
            return [], []

        has_header = False
        try:
            has_header = csv.Sniffer().has_header(sample)
        except Exception:
            has_header = False

        if has_header:
            headers = [self._normalize_header(h) for h in rows[0]]
            data_rows = rows[1:]
        else:
            headers = [f"column_{i + 1}" for i in range(len(rows[0]))]
            data_rows = rows

        normalized_rows = []
        for row in data_rows:
            normalized_rows.append([col.strip() if isinstance(col, str) else col for col in row])

        return normalized_rows, headers

    def _parse_excel(self):
        decoded = self._decode_binary_file()
        filename = (self.file_name or "").lower()

        if filename.endswith(".xlsx"):
            try:
                import openpyxl
            except Exception as exc:
                raise UserError(
                    "openpyxl is required to read .xlsx files. Please install it on the server."
                ) from exc

            wb = openpyxl.load_workbook(io.BytesIO(decoded), read_only=True, data_only=True)
            requested = (self.sheet_name or "").strip()
            if requested:
                if requested in wb.sheetnames:
                    sheet = wb[requested]
                else:
                    available = ", ".join(wb.sheetnames)
                    wb.close()
                    raise UserError(
                        f"Sheet '{requested}' was not found in this workbook. Available sheets: {available}"
                    )
            else:
                sheet = wb[wb.sheetnames[-1]]
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([col if col is not None else "" for col in row])
            wb.close()
        else:
            try:
                import xlrd
            except Exception as exc:
                raise UserError(
                    "xlrd is required to read .xls files. Please install it on the server."
                ) from exc

            try:
                wb = xlrd.open_workbook(file_contents=decoded)
                requested = (self.sheet_name or "").strip()
                if requested and requested in wb.sheet_names():
                    sheet = wb.sheet_by_name(requested)
                else:
                    sheet = wb.sheet_by_index(wb.nsheets - 1)
                rows = []
                for r in range(sheet.nrows):
                    rows.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
            except Exception as exc:
                if self._looks_like_html_table(decoded):
                    rows = self._parse_html_table(decoded)
                else:
                    raise UserError(
                        "This .xls file does not appear to be a valid Excel file. "
                        "Please re-save it as a real .xls/.xlsx, or export as .csv."
                    ) from exc

        rows = [row for row in rows if any(str(col).strip() for col in row)]
        if not rows:
            return [], []

        headers = [self._normalize_header(str(h)) for h in rows[0]]
        data_rows = rows[1:]

        normalized_rows = []
        for row in data_rows:
            normalized_rows.append([str(col).strip() if col is not None else "" for col in row])
        return normalized_rows, headers

    def _looks_like_html_table(self, decoded_bytes):
        head = (decoded_bytes or b"")[:4096].lstrip().lower()
        if not head.startswith(b"<"):
            return False
        return b"<table" in head or b"<html" in head

    def _parse_html_table(self, decoded_bytes):
        try:
            from lxml import html
        except Exception as exc:
            raise UserError(
                "This file looks like an HTML table saved as .xls, but lxml is not available "
                "to parse it. Please export as CSV or save as a real Excel file."
            ) from exc

        text = (decoded_bytes or b"").decode("utf-8", errors="ignore")
        doc = html.fromstring(text)
        tables = doc.xpath("//table")
        if not tables:
            return []

        rows = []
        for tr in tables[0].xpath(".//tr"):
            cells = tr.xpath("./th|./td")
            row = []
            for cell in cells:
                value = "".join(cell.itertext()).strip()
                row.append(value)
            if any(v for v in row):
                rows.append(row)
        return rows

    def _decode_file_data(self):
        decoded = base64.b64decode(self.file_data or b"")
        try:
            return decoded.decode("utf-8")
        except UnicodeDecodeError:
            return decoded.decode("latin1")

    def _decode_binary_file(self):
        return base64.b64decode(self.file_data or b"")

    def _normalize_header(self, header):
        header = (header or "").strip().lower()
        header = re.sub(r"[\s\-_]+", " ", header)
        return header

    def _detect_qty_key(self, headers):
        if not headers:
            return None

        normalized = [(h or "").strip().lower() for h in headers]

        strong = ("qty", "quantity", "qtty", "count", "pcs", "pieces")
        for header in normalized:
            if header in strong or any(header.startswith(t) for t in strong):
                return header

        for header in normalized:
            if header in ("q", "qty.", "qtty."):
                return header

        weak = ("qty", "quantity", "qtty", "count")
        for header in normalized:
            if any(token in header for token in weak) and not any(
                token in header for token in ("amount", "total", "value", "price", "rate")
            ):
                return header

        for header in normalized:
            if header in ("s", "sales", "sold", "pcs", "pc"):
                return header

        return None

    def _detect_descriptor_keys(self, headers, qty_key):
        ignore_tokens = (
            "qty",
            "quantity",
            "qtty",
            "amount",
            "count",
            "price",
            "unit price",
            "rate",
            "total",
            "value",
        )
        keys = []
        for header in headers:
            if header == qty_key:
                continue
            if any(token in header for token in ignore_tokens):
                continue
            keys.append(header)
        if not keys:
            keys = [h for h in headers if h != qty_key]
        return keys

    def _build_descriptor(self, row, headers, descriptor_keys, qty_key):
        row_map = dict(zip(headers, row))
        parts = []
        for key in descriptor_keys:
            value = row_map.get(key)
            if value is None:
                continue
            value = str(value).strip()
            if not value:
                continue
            if re.fullmatch(r"\d+(?:\.\d+)?", value):
                continue
            parts.append(value)
        return " ".join(parts).strip()

    def _extract_qty(self, row, headers, qty_key):
        if not qty_key:
            return 1.0
        row_map = dict(zip(headers, row))
        raw = row_map.get(qty_key)
        if raw is None or raw == "":
            return 1.0
        try:
            return float(raw)
        except Exception:
            raw = str(raw).strip()
            match = re.search(r"(\d+(?:\.\d+)?)", raw)
            if match:
                return float(match.group(1))
        return 1.0

    def _render_row(self, row, headers):
        row_map = dict(zip(headers, row))
        rendered = []
        for header in headers:
            value = row_map.get(header)
            if value in (None, ""):
                continue
            rendered.append(f"{header}: {value}")
        return "; ".join(rendered)

    def _match_product(self, row, headers, descriptor):
        exact = self._find_exact_match(row, headers)
        if exact:
            return {"status": "matched", "product_id": exact.id, "candidate_ids": [exact.id]}

        candidates = self._search_candidates(descriptor)
        if not candidates:
            return {"status": "no_match", "candidate_ids": []}

        descriptor_size = self._extract_size_token(descriptor)
        scored = []
        for product in candidates:
            scored.append((self._similarity_score(descriptor, product, descriptor_size), product))

        scored.sort(key=lambda s: s[0], reverse=True)
        top_score, top_product = scored[0]
        candidate_ids = [p.id for _, p in scored[:5]]

        second_score = scored[1][0] if len(scored) > 1 else 0.0
        if top_score >= 0.75 and (top_score - second_score) >= 0.1:
            return {"status": "matched", "product_id": top_product.id, "candidate_ids": candidate_ids}

        if top_score >= 0.6:
            return {"status": "multiple", "candidate_ids": candidate_ids}

        return {"status": "no_match", "candidate_ids": candidate_ids}

    def _find_exact_match(self, row, headers):
        row_map = dict(zip(headers, row))
        exact_tokens = ("sku", "code", "barcode", "item code", "product code")
        for header, value in row_map.items():
            if not value:
                continue
            if any(token in header for token in exact_tokens):
                term = str(value).strip()
                if not term:
                    continue
                product = self.env["product.product"].search(
                    [
                        ("type", "in", ("product", "consu")),
                        "|",
                        ("default_code", "=", term),
                        ("barcode", "=", term),
                    ],
                    limit=1,
                )
                if product:
                    return product
        return False

    def _search_candidates(self, descriptor):
        Product = self.env["product.product"]
        if not descriptor:
            return Product

        descriptor = descriptor.strip()
        tokens = [t for t in re.split(r"[\W_]+", descriptor) if len(t) >= 3]

        domain = ["&", ("type", "in", ("product", "consu")), "|", "|",
                  ("default_code", "ilike", descriptor),
                  ("barcode", "ilike", descriptor),
                  ("name", "ilike", descriptor)]
        candidates = Product.search(domain, limit=40)
        if candidates:
            return candidates

        if tokens:
            token = tokens[0]
            domain = ["&", ("type", "in", ("product", "consu")), "|",
                      ("default_code", "ilike", token),
                      ("name", "ilike", token)]
            candidates = Product.search(domain, limit=40)
            if candidates:
                return candidates

        return Product.search([("type", "in", ("product", "consu"))], limit=10)

    def _similarity_score(self, descriptor, product, descriptor_size=None):
        text = (descriptor or "").lower().strip()
        if not text:
            return 0.0

        candidates = [
            product.display_name,
            product.name,
            product.default_code or "",
            product.barcode or "",
        ]
        score = 0.0
        for candidate in candidates:
            if not candidate:
                continue
            ratio = SequenceMatcher(None, text, candidate.lower()).ratio()
            score = max(score, ratio)
            if descriptor_size:
                candidate_size = self._extract_size_token(candidate)
                if candidate_size and candidate_size == descriptor_size:
                    score = max(score, min(1.0, ratio + 0.12))
        return score

    def _extract_size_token(self, text):
        if not text:
            return None
        text = str(text).lower()
        text = re.sub(r"(\d)\s+(ml|g|kg|oz|l|litre|liter|lb|lbs)", r"\1\2", text)
        patterns = [
            r"\b(\d+(?:\.\d+)?)(ml|g|kg|oz|l|litre|liter|lb|lbs)\b",
            r"\b(\d+)(?:x|Ã—)(\d+)(ml|g|kg|oz|l|litre|liter|lb|lbs)\b",
        ]
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                return "".join(match.groups())
        return None


class StockPickingBatchImportLine(models.TransientModel):
    _name = "stock.picking.batch.import.line"
    _description = "Stock Picking Batch Import Line"

    wizard_id = fields.Many2one(
        "stock.picking.batch.import.wizard", required=True, ondelete="cascade"
    )
    row_no = fields.Integer(string="Row")
    raw_data = fields.Text(string="Raw Data")
    description = fields.Char(string="Description")
    quantity = fields.Float(string="Quantity", default=1.0)
    match_status = fields.Selection(
        [
            ("matched", "Matched"),
            ("multiple", "Multiple Matches"),
            ("no_match", "No Match"),
        ],
        string="Status",
    )
    product_id = fields.Many2one(
        "product.product",
        string="Product",
        domain=[("type", "in", ("product", "consu"))],
    )
    candidate_product_ids = fields.Many2many("product.product", string="Candidates")
