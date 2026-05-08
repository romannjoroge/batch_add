# -*- coding: utf-8 -*-
import base64
import csv
import io
import re
from difflib import SequenceMatcher

from odoo import api, fields, models
from odoo.exceptions import UserError


class BatchImportWizardMixin(models.AbstractModel):
    _name = "batch_add.batch_import_wizard_mixin"
    _description = "Batch Import Wizard Mixin"

    file_data = fields.Binary(string="CSV File", required=True)
    file_name = fields.Char(string="File Name")
    sheet_name = fields.Char(
        string="Sheet",
        help="For Excel files with multiple sheets, pick the sheet to import from.",
    )

    @api.onchange("file_data", "file_name")
    def _onchange_file_data_set_sheet(self):
        for wizard in self:
            if not wizard.file_name or not wizard.file_data:
                wizard.sheet_name = False
                continue
            filename = (wizard.file_name or "").lower()
            if not filename.endswith((".xlsx", ".xls")):
                wizard.sheet_name = False
                continue

            sheet_names = wizard._get_excel_sheet_names()
            wizard.sheet_name = sheet_names[-1] if sheet_names else False

    def action_parse_file(self):
        self.ensure_one()
        self._validate_file_input()
        self.line_ids.unlink()

        rows, headers = self._parse_file()
        if not rows:
            raise UserError("No rows were found in the file.")

        qty_key = self._detect_qty_key(headers)
        descriptor_keys = self._detect_descriptor_keys(headers, qty_key)

        preview_lines = []
        for index, row in enumerate(rows, start=1):
            descriptor = self._build_descriptor(row, headers, descriptor_keys, qty_key)
            quantity = self._extract_qty(row, headers, qty_key)
            if not descriptor:
                continue
            match = self._match_product(row, headers, descriptor)
            preview_lines.append(
                (
                    0,
                    0,
                    {
                        "row_no": index,
                        "raw_data": self._render_row(row, headers),
                        "description": descriptor,
                        "quantity": quantity,
                        "match_status": match["status"],
                        "product_id": match.get("product_id"),
                        "candidate_product_ids": [(6, 0, match.get("candidate_ids", []))],
                    },
                )
            )

        self.line_ids = preview_lines
        return self._get_reopen_action()

    def _get_reopen_action(self):
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "views": [(False, "form")],
            "target": "new",
        }

    def _validate_file_input(self):
        if not self.file_data:
            raise UserError("Please upload a CSV file first.")
        if self.file_name and not self._is_supported_file(self.file_name):
            raise UserError("Please upload a CSV or Excel file (.csv, .xls, .xlsx).")

    def _get_excel_sheet_names(self):
        decoded = self._decode_binary_file()
        filename = (self.file_name or "").lower()

        if filename.endswith(".xlsx"):
            try:
                import openpyxl
            except Exception:
                return []
            workbook = openpyxl.load_workbook(
                io.BytesIO(decoded), read_only=True, data_only=True
            )
            sheet_names = list(workbook.sheetnames or [])
            workbook.close()
            return sheet_names

        if filename.endswith(".xls"):
            if self._looks_like_html_table(decoded):
                return ["Sheet1"]
            try:
                import xlrd
            except Exception:
                return []
            workbook = xlrd.open_workbook(file_contents=decoded)
            return list(getattr(workbook, "sheet_names", lambda: [])())

        return []

    def _is_supported_file(self, filename):
        return (filename or "").lower().endswith((".csv", ".xls", ".xlsx"))

    def _parse_file(self):
        filename = (self.file_name or "").lower()
        if filename.endswith((".xls", ".xlsx")):
            return self._parse_excel()
        return self._parse_csv()

    def _parse_csv(self):
        decoded = self._decode_file_data()
        sample = decoded[:4096]

        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=[",", ";", "\t", "|"])
        except Exception:
            dialect = csv.get_dialect("excel")

        reader = csv.reader(io.StringIO(decoded), dialect)
        rows = list(reader)
        rows = [row for row in rows if any(col.strip() for col in row)]
        if not rows:
            return [], []

        try:
            has_header = csv.Sniffer().has_header(sample)
        except Exception:
            has_header = False

        if has_header:
            headers = [self._normalize_header(h) for h in rows[0]]
            data_rows = rows[1:]
        else:
            headers = [f"column_{i + 1}" for i in range(len(rows[0]))]
            data_rows = rows

        normalized_rows = []
        for row in data_rows:
            normalized_rows.append(
                [col.strip() if isinstance(col, str) else col for col in row]
            )
        return normalized_rows, headers

    def _parse_excel(self):
        decoded = self._decode_binary_file()
        filename = (self.file_name or "").lower()

        if filename.endswith(".xlsx"):
            try:
                import openpyxl
            except Exception as exc:
                raise UserError(
                    "openpyxl is required to read .xlsx files. Please install it on the server."
                ) from exc

            workbook = openpyxl.load_workbook(
                io.BytesIO(decoded), read_only=True, data_only=True
            )
            requested = (self.sheet_name or "").strip()
            if requested:
                if requested in workbook.sheetnames:
                    sheet = workbook[requested]
                else:
                    available = ", ".join(workbook.sheetnames)
                    workbook.close()
                    raise UserError(
                        f"Sheet '{requested}' was not found in this workbook. Available sheets: {available}"
                    )
            else:
                sheet = workbook[workbook.sheetnames[-1]]

            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([col if col is not None else "" for col in row])
            workbook.close()
        else:
            try:
                import xlrd
            except Exception as exc:
                raise UserError(
                    "xlrd is required to read .xls files. Please install it on the server."
                ) from exc

            try:
                workbook = xlrd.open_workbook(file_contents=decoded)
                requested = (self.sheet_name or "").strip()
                if requested and requested in workbook.sheet_names():
                    sheet = workbook.sheet_by_name(requested)
                else:
                    sheet = workbook.sheet_by_index(workbook.nsheets - 1)
                rows = []
                for row_index in range(sheet.nrows):
                    rows.append(
                        [sheet.cell_value(row_index, col_index) for col_index in range(sheet.ncols)]
                    )
            except Exception as exc:
                if self._looks_like_html_table(decoded):
                    rows = self._parse_html_table(decoded)
                else:
                    raise UserError(
                        "This .xls file does not appear to be a valid Excel file. "
                        "Please re-save it as a real .xls/.xlsx, or export as .csv."
                    ) from exc

        rows = [row for row in rows if any(str(col).strip() for col in row)]
        if not rows:
            return [], []

        headers = [self._normalize_header(str(header)) for header in rows[0]]
        data_rows = rows[1:]
        normalized_rows = []
        for row in data_rows:
            normalized_rows.append(
                [str(col).strip() if col is not None else "" for col in row]
            )
        return normalized_rows, headers

    def _looks_like_html_table(self, decoded_bytes):
        head = (decoded_bytes or b"")[:4096].lstrip().lower()
        return head.startswith(b"<") and (b"<table" in head or b"<html" in head)

    def _parse_html_table(self, decoded_bytes):
        try:
            from lxml import html
        except Exception as exc:
            raise UserError(
                "This file looks like an HTML table saved as .xls, but lxml is not available "
                "to parse it. Please export as CSV or save as a real Excel file."
            ) from exc

        text = (decoded_bytes or b"").decode("utf-8", errors="ignore")
        document = html.fromstring(text)
        tables = document.xpath("//table")
        if not tables:
            return []

        rows = []
        for tr in tables[0].xpath(".//tr"):
            cells = tr.xpath("./th|./td")
            row = []
            for cell in cells:
                row.append("".join(cell.itertext()).strip())
            if any(row):
                rows.append(row)
        return rows

    def _decode_file_data(self):
        decoded = base64.b64decode(self.file_data or b"")
        for encoding in ("utf-8-sig", "utf-8", "latin1"):
            try:
                return decoded.decode(encoding)
            except UnicodeDecodeError:
                continue
        return decoded.decode("latin1", errors="ignore")

    def _decode_binary_file(self):
        return base64.b64decode(self.file_data or b"")

    def _normalize_header(self, header):
        normalized = (header or "").strip().lower().replace("\ufeff", "")
        return re.sub(r"[\s\-_]+", " ", normalized)

    def _detect_qty_key(self, headers):
        if not headers:
            return None

        normalized = [(header or "").strip().lower() for header in headers]
        strong_tokens = ("qty", "quantity", "qtty", "count", "pcs", "pieces")
        for header in normalized:
            if header in strong_tokens or any(header.startswith(token) for token in strong_tokens):
                return header

        for header in normalized:
            if header in ("q", "qty.", "qtty."):
                return header

        weak_tokens = ("qty", "quantity", "qtty", "count")
        monetary_tokens = ("amount", "total", "value", "price", "rate")
        for header in normalized:
            if any(token in header for token in weak_tokens) and not any(
                token in header for token in monetary_tokens
            ):
                return header

        for header in normalized:
            if header in ("s", "sales", "sold", "pcs", "pc"):
                return header

        return None

    def _detect_descriptor_keys(self, headers, qty_key):
        ignore_tokens = (
            "qty",
            "quantity",
            "qtty",
            "amount",
            "count",
            "price",
            "unit price",
            "rate",
            "total",
            "value",
        )
        keys = []
        for header in headers:
            if header == qty_key:
                continue
            if any(token in header for token in ignore_tokens):
                continue
            keys.append(header)
        return keys or [header for header in headers if header != qty_key]

    def _build_descriptor(self, row, headers, descriptor_keys, qty_key):
        row_map = dict(zip(headers, row))
        parts = []
        for key in descriptor_keys:
            value = row_map.get(key)
            if value is None:
                continue
            value = str(value).strip()
            if not value or re.fullmatch(r"\d+(?:\.\d+)?", value):
                continue
            parts.append(value)
        return " ".join(parts).strip()

    def _extract_qty(self, row, headers, qty_key):
        if not qty_key:
            return 1.0
        row_map = dict(zip(headers, row))
        raw = row_map.get(qty_key)
        if raw in (None, ""):
            return 1.0
        try:
            return float(raw)
        except Exception:
            match = re.search(r"(\d+(?:\.\d+)?)", str(raw).strip())
            if match:
                return float(match.group(1))
        return 1.0

    def _render_row(self, row, headers):
        row_map = dict(zip(headers, row))
        rendered = []
        for header in headers:
            value = row_map.get(header)
            if value in (None, ""):
                continue
            rendered.append(f"{header}: {value}")
        return "; ".join(rendered)

    def _match_product(self, row, headers, descriptor):
        exact_match = self._find_exact_match(row, headers)
        if exact_match:
            return {
                "status": "matched",
                "product_id": exact_match.id,
                "candidate_ids": [exact_match.id],
            }

        candidates = self._search_candidates(descriptor)
        if not candidates:
            return {"status": "no_match", "candidate_ids": []}

        descriptor_size = self._extract_size_token(descriptor)
        scored = [
            (self._similarity_score(descriptor, product, descriptor_size), product)
            for product in candidates
        ]
        scored.sort(key=lambda item: item[0], reverse=True)

        top_score, top_product = scored[0]
        candidate_ids = [product.id for _, product in scored[:5]]
        second_score = scored[1][0] if len(scored) > 1 else 0.0

        if top_score >= 0.75 and (top_score - second_score) >= 0.1:
            return {
                "status": "matched",
                "product_id": top_product.id,
                "candidate_ids": candidate_ids,
            }
        if top_score >= 0.6:
            return {"status": "multiple", "candidate_ids": candidate_ids}
        return {"status": "no_match", "candidate_ids": candidate_ids}

    def _find_exact_match(self, row, headers):
        row_map = dict(zip(headers, row))
        exact_tokens = ("sku", "code", "barcode", "item code", "product code")
        for header, value in row_map.items():
            if not value or not any(token in header for token in exact_tokens):
                continue
            term = str(value).strip()
            if not term:
                continue
            product = self.env["product.product"].search(
                self._get_exact_match_domain(term),
                limit=1,
            )
            if product:
                return product
        return False

    def _get_exact_match_domain(self, term):
        return self._get_product_domain() + [
            "|",
            ("default_code", "=", term),
            ("barcode", "=", term),
        ]

    def _search_candidates(self, descriptor):
        product_model = self.env["product.product"]
        base_domain = self._get_product_domain()
        if not descriptor:
            return product_model.search(base_domain, limit=10)

        descriptor = descriptor.strip()
        tokens = [token for token in re.split(r"[\W_]+", descriptor) if len(token) >= 3]

        candidates = product_model.search(
            base_domain
            + [
                "|",
                "|",
                ("default_code", "ilike", descriptor),
                ("barcode", "ilike", descriptor),
                ("name", "ilike", descriptor),
            ],
            limit=40,
        )
        if candidates:
            return candidates

        if tokens:
            token = tokens[0]
            candidates = product_model.search(
                base_domain + ["|", ("default_code", "ilike", token), ("name", "ilike", token)],
                limit=40,
            )
            if candidates:
                return candidates

        return product_model.search(base_domain, limit=10)

    def _similarity_score(self, descriptor, product, descriptor_size=None):
        text = (descriptor or "").lower().strip()
        if not text:
            return 0.0

        score = 0.0
        for candidate in (
            product.display_name,
            product.name,
            product.default_code or "",
            product.barcode or "",
        ):
            if not candidate:
                continue
            ratio = SequenceMatcher(None, text, candidate.lower()).ratio()
            score = max(score, ratio)
            if descriptor_size:
                candidate_size = self._extract_size_token(candidate)
                if candidate_size and candidate_size == descriptor_size:
                    score = max(score, min(1.0, ratio + 0.12))
        return score

    def _extract_size_token(self, text):
        if not text:
            return None
        text = str(text).lower()
        text = re.sub(r"(\d)\s+(ml|g|kg|oz|l|litre|liter|lb|lbs)", r"\1\2", text)
        patterns = [
            r"\b(\d+(?:\.\d+)?)(ml|g|kg|oz|l|litre|liter|lb|lbs)\b",
            r"\b(\d+)(?:x|\u00d7|\*)(\d+)(ml|g|kg|oz|l|litre|liter|lb|lbs)\b",
        ]
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                return "".join(match.groups())
        return None

    def _get_product_domain(self):
        raise NotImplementedError("Subclasses must define a product domain.")
